import os
import pytesseract
from PIL import Image, ImageEnhance
from src.config import DEFAULT_TESSERACT_PATHS

# Auto-configure pytesseract cmd path
for t_path in DEFAULT_TESSERACT_PATHS:
    if os.path.exists(t_path):
        pytesseract.pytesseract.tesseract_cmd = t_path
        break

def preprocess_image(image: Image.Image) -> Image.Image:
    """Preprocess image for better OCR accuracy."""
    try:
        gray = image.convert('L')
        enhancer = ImageEnhance.Contrast(gray)
        enhanced = enhancer.enhance(1.8)
        return enhanced
    except Exception:
        return image

def get_tesseract_lang() -> str:
    """Detect available languages in Tesseract to prevent errors."""
    try:
        langs = pytesseract.get_languages()
        if 'ind' in langs and 'eng' in langs:
            return 'ind+eng'
        elif 'ind' in langs:
            return 'ind'
        elif 'eng' in langs:
            return 'eng'
    except Exception:
        pass
    return ''

def extract_text_from_docx(file_path: str) -> str:
    """Extract text from Word .docx file including paragraphs and tables."""
    try:
        import docx
        doc = docx.Document(file_path)
        parts = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                parts.append(text)
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    deduped = []
                    for t in row_texts:
                        if not deduped or deduped[-1] != t:
                            deduped.append(t)
                    parts.append(" | ".join(deduped))
        return "\n\n".join(parts).strip()
    except Exception as e:
        raise ValueError(f"Failed to read DOCX file: {e}")

def extract_text_from_image(image_path: str) -> str:
    """Extract text from an image file using Tesseract OCR."""
    if not os.path.exists(image_path):
        raise FileNotFoundError(f"Image file not found: {image_path}")
    
    img = Image.open(image_path)
    processed_img = preprocess_image(img)
    
    lang = get_tesseract_lang()
    try:
        if lang:
            text = pytesseract.image_to_string(processed_img, lang=lang)
        else:
            text = pytesseract.image_to_string(processed_img)
    except Exception:
        text = pytesseract.image_to_string(processed_img)
        
    return text.strip()

def extract_text_from_file(file_path: str) -> str:
    """Extract text from supported document/image files (PNG, JPG, WEBP, TXT, MD, PDF, DOCX)."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
    
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext in ['.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tiff', '.tif']:
        return extract_text_from_image(file_path)
    
    elif ext in ['.txt', '.md', '.log', '.csv', '.json']:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read().strip()

    elif ext in ['.docx', '.doc']:
        return extract_text_from_docx(file_path)
            
    elif ext == '.pdf':
        try:
            import fitz  # PyMuPDF
            import io
            doc = fitz.open(file_path)
            extracted_pages = []
            for page_num, page in enumerate(doc):
                text = page.get_text().strip()
                if not text:
                    # Scanned PDF page fallback: rasterize page to image and use Tesseract OCR
                    try:
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = io.BytesIO(pix.tobytes("png"))
                        page_img = Image.open(img_bytes)
                        processed_img = preprocess_image(page_img)
                        tess_config = r'--oem 3 --psm 6'
                        lang = get_tesseract_lang()
                        try:
                            if lang:
                                text = pytesseract.image_to_string(processed_img, lang=lang, config=tess_config).strip()
                            else:
                                text = pytesseract.image_to_string(processed_img, config=tess_config).strip()
                        except Exception:
                            text = pytesseract.image_to_string(processed_img).strip()
                    except Exception as ocr_err:
                        print(f"[OCR Warning] Scanned PDF page {page_num+1} OCR fallback error: {ocr_err}")
                        text = ""
                if text:
                    extracted_pages.append(text)
            return "\n\n".join(extracted_pages).strip()
        except ImportError:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read().strip()
    else:
        raise ValueError(f"Unsupported file format: {ext}. Supported: PNG, JPG, WEBP, TXT, MD, PDF, DOCX")

